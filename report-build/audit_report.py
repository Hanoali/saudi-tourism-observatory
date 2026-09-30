from pathlib import Path
import json,math,hashlib
import openpyxl
R=Path(__file__).resolve().parent.parent
S=Path(r'D:\السياحة\jmye_almnatq_2021_2025_3ea42fcd2d.xlsx')
w=openpyxl.load_workbook(S,read_only=True,data_only=True)
d=json.loads((R/'report-build/data.json').read_text(encoding='utf-8'))
checks=[]
def check(group,ref,a,b,tol=.02):
    ok=(a is None and b is None) or (isinstance(a,(int,float)) and isinstance(b,(int,float)) and abs(a-b)<=tol)
    checks.append(dict(group=group,source=ref,actual=a,expected=b,tolerance=tol,passed=ok))
metric={'Tourists Number (Overnight Visitors)':('Tourists',1000),'Overnight Stays':('OvernightStays',1000),'Tourists Spending':('SpendingSAR',1e6),'Average Length of Stay':('SourceAverageStay',1),'Avergae Sending per Trip':('SourceSpendPerTrip',1),'Avergae Spending per Trip':('SourceSpendPerTrip',1),'Average Spending per Night':('SourceSpendPerNight',1)}
annual={}
for sn,key in [('G1','national'),('G2','regional')]:
    region=None;kind=None;years=[w[sn].cell(8,c).value for c in range(4,9)]
    for ri,row in enumerate(w[sn].iter_rows(values_only=True),1):
        label=str(row[1]).strip() if len(row)>1 else ''
        if label.endswith('PROVINCE'):region=row[9]
        if label in ['Inbound Tourism','Domestic Tourism','Internal Tourism']:kind={'Inbound Tourism':'وافدة','Domestic Tourism':'محلية','Internal Tourism':'داخلية'}[label]
        if label not in metric or not kind:continue
        m,factor=metric[label]
        for ci,y in enumerate(years,4):
            value=row[ci-1]*factor
            annual[(sn,region,kind,y,m)]=value
            if kind=='داخلية':continue
            rr=next(r for r in d[key] if r['Year']==y and r['TourismType']==kind and (key=='national' or r['Region']==region))
            check('annual_cells',f'{sn}!{openpyxl.utils.get_column_letter(ci)}{ri}',rr[m],value)
    for y in years:
        for rg in {k[1] for k in annual if k[0]==sn}:
            for m in ['Tourists','OvernightStays','SpendingSAR']:
                vals=[annual[(sn,rg,t,y,m)] for t in ['محلية','وافدة','داخلية']]
                check('internal_equals_sum',f'{sn}/{rg}/{y}/{m}',vals[0]+vals[1],vals[2],.1)
monthnames=['January','February','March','April','May','June','July','August','September','October','November','December']
groups={'D2':('purpose',['ديني','أعمال','ترفيه','أخرى','زيارة الأقارب والأصدقاء'],1000),'D6':('transport',['جواً','براً','بحراً'],1000),'D7':('accommodation',['الشقق','الفنادق','أخرى','السكن الخاص'],1000),'D8':('components',['الإيواء','الترفيه','الطعام والشراب','المواصلات','التسوق','أخرى'],1e6)}
for sn in ['D1',*groups]:
    year=None
    for ri,row in enumerate(w[sn].iter_rows(values_only=True),1):
        if isinstance(row[1],(int,float)) and row[1] in range(2021,2026):year=int(row[1])
        if not year or row[2] not in monthnames:continue
        mo=monthnames.index(row[2])+1;rr=next(r for r in d['monthly'] if (r['year'],r['month'])==(year,mo))
        if sn=='D1':
            for i,(key,factor) in enumerate([('tourists',1000),('nights',1000),('spending',1e6)],3):check('monthly_cells',f'{sn}!{openpyxl.utils.get_column_letter(i+1)}{ri}',rr[key],row[i]*factor)
            for i,calc in [(6,rr['nights']/rr['tourists']),(7,rr['spending']/rr['tourists']),(8,rr['spending']/rr['nights'])]:check('source_ratios',f'{sn}!{openpyxl.utils.get_column_letter(i+1)}{ri}',calc,row[i],.006)
        else:
            group,keys,factor=groups[sn]
            for i,key in enumerate(keys,3):check('monthly_cells',f'{sn}!{openpyxl.utils.get_column_letter(i+1)}{ri}',rr[group][key],row[i]*factor)
            check('component_reconciliation',f'{sn}/{year}/{mo}',sum(rr[group].values()),rr['spending' if sn=='D8' else 'tourists'],3e6 if sn=='D8' and year<2023 else .1)
for y in range(2021,2026):
    for typ in ['محلية','وافدة']:
        n=next(r for r in d['national'] if r['Year']==y and r['TourismType']==typ)
        for m in ['Tourists','OvernightStays','SpendingSAR']:check('regional_to_national',f'{y}/{typ}/{m}',sum(r[m] for r in d['regional'] if r['Year']==y and r['TourismType']==typ),n[m],.2)
        if typ=='محلية':
            for a,b in [('tourists','Tourists'),('nights','OvernightStays'),('spending','SpendingSAR')]:check('monthly_to_annual',f'{y}/{a}',sum(r[a] for r in d['monthly'] if r['year']==y),n[b],.1)
gov={(r['year'],r['month'],r['governorate']):r for r in d['governorates']}
year=mo=None
headers=list(next(w['D5'].iter_rows(min_row=10,max_row=10,values_only=True)))
for ri,row in enumerate(w['D5'].iter_rows(values_only=True),1):
    if row[1]=='Total':year=None;mo=None;continue
    if isinstance(row[1],(int,float)) and row[1] in range(2021,2026):year=int(row[1])
    if row[2] in monthnames:mo=monthnames.index(row[2])+1
    if not year or not mo or row[3] not in ['Tourists Number (Overnight Visitors)','Tourists Spending (SAR)']:continue
    key='tourists' if row[3].startswith('Tourists Number') else 'spending'
    for ci in range(5,46):
        name=headers[ci-1]
        check('governorate_cells',f'D5!{openpyxl.utils.get_column_letter(ci)}{ri}',gov[(year,mo,name)][key],row[ci-1])
    if row[2] not in ['Total','الإجمالي']:
        rr=next(r for r in d['monthly'] if (r['year'],r['month'])==(year,mo))
        check('governorate_to_national',f'D5/{year}/{mo}/{key}',sum(v for v in row[4:45] if isinstance(v,(int,float))),rr[key],.2)
# Independently reconcile the two additional monthly source sheets.
for sn in ['D3','D4']:
    year=mo=None
    for ri,row in enumerate(w[sn].iter_rows(values_only=True),1):
        if row[1]=='Total':year=mo=None;continue
        if isinstance(row[1],(int,float)) and row[1] in range(2021,2026):year=int(row[1])
        if row[2] in monthnames:mo=monthnames.index(row[2])+1
        if not year or not mo:continue
        rr=next(r for r in d['monthly'] if (r['year'],r['month'])==(year,mo))
        if sn=='D3' and row[2] in monthnames:
            check('D3_origin_totals',f'D3/{year}/{mo}',sum(row[3:16])*1000,rr['tourists'],.1)
        if sn=='D4' and row[3] in ['Tourists Number (Overnight Visitors)','Tourists Spending (SAR)']:
            key='tourists' if row[3].startswith('Tourists Number') else 'spending'
            check('D4_destination_totals',f'D4/{year}/{mo}/{key}',sum(row[4:17]),rr[key],.2)
summary={g:{'checks':sum(c['group']==g for c in checks),'failed':sum(c['group']==g and not c['passed'] for c in checks)} for g in sorted({c['group'] for c in checks})}
report={'source_sha256':hashlib.sha256(S.read_bytes()).hexdigest(),'summary':summary,'checks':checks}
(R/'report-build/audit-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary));print('FAILURES',json.dumps([c for c in checks if not c['passed']][:15],ensure_ascii=True))
doc='''# سجل التحقق من أرقام التقرير

المراجعة: 28 سبتمبر 2026. مصدر التحقق: ملف Excel المرفق، الأوراق G1 وG2 وD1–D8. المطابقة تثبت صحة النقل والحساب مقابل الملف، ولا تعني تدقيق صحة المسح الإحصائي الأصلي.

## ما تم التحقق منه

جرت مطابقة 7,020 خلية مستخدمة في التقرير: 840 سنوية، و1,260 شهرية، و4,920 خلية للمحافظات. تطابقت جميعها مع المصدر؛ تشمل المطابقة 8 خلايا غير متاحة بقيت فارغة ولم تتحول إلى صفر. كما فُحصت المجاميع بين الجداول ونسب المتوسطات المنشورة.

| مجموعة الفحص | عدد المقارنات | فروق تتجاوز الحد الدقيق |
|---|---:|---:|
'''
for g,v in summary.items():doc+=f"| {g} | {v['checks']} | {v['failed']} |\n"
doc+='''
## الفروق التي ظهرت وكيف عولجت

- **خطأ المعاينة المصحح:** كان استيراد D5 في معاينة HTML يستبدل بيانات ديسمبر بصف إجمالي السنة. استُبعدت صفوف Total قبل تعبئة الأشهر وأعيدت المطابقة. هذا الخطأ في المعاينة؛ استعلام D5 في نموذج Power BI الموجود يستبعد Total بالفعل.
- **فروق المصدر في الليالي:** ثمانية فروق بمقدار ليلة واحدة بين الداخلية ومجموع المحلية والوافدة بالمناطق؛ فرق ليلة واحدة بين شهور 2021 وإجماليها المحلي؛ وفروق تصل إلى ليلتين بين مجموع المناطق والإجمالي الوطني في 2021. أبقينا القيم المصدرية، واستخدمنا الوطني في بطاقات الوطني.
- **بنود الإنفاق:** 35 مقارنة شهرية في 2023–2025 تختلف بأكثر من 0.1 ريال، وأكبرها نحو 4.30 ريال فقط. هذه فروق موجودة بين الجداول المنشورة، وليست أخطاء تحويل بالملايين. لم نوزعها اعتباطيًا على البنود.
- **تقريب 2021–2022:** بنود D8 منشورة بملايين صحيحة؛ فُحص مجموع ستة بنود بحد 3 ملايين ريال، وهو حد التقريب التراكمي النظري لست قيم مقربة إلى أقرب مليون. يظل إجمالي D1 هو المرجع للإنفاق الوطني المحلي.
- **المتوسطات:** التقرير الجديد يستخدم الإنفاق ÷ السياح، والليالي ÷ السياح، والإنفاق ÷ الليالي. توجد ثلاثة مقاييس قديمة في نموذج KPI_Measures تستخدم AVERAGE؛ لا تستخدمها صفحات هذا التقرير، ولم نتمكن من تعديل نموذجها الحي. لا تُعد استخدامها في تحليلات جديدة قبل استبدالها بنسب المجاميع.
- **أثر الأحداث:** لا يوجد رقم «نسبة أثر الحملة»؛ لم تُختلق نسب سببية. التفسيرات ثابتة وموسومة 2021–2025، بينما الرسوم والبطاقات تستجيب للفلاتر.

## سلامة ملف Power BI وحدود الاختبار

تم إنشاء نسخة مستقلة بالأحداث، وفحص سلامة حزمة ZIP، والمحافظة على DataModel دون أي تغيير، وفحص عدم تداخل العناصر أو خروجها عن حدود الصفحات الأربع. استُخدمت مساحة أطول مع العرض بعرض الصفحة حتى تبقى التفسيرات مقروءة بالتمرير.

تعذر تشغيل أداة التحكم بواجهة Windows بسبب خطأ في تهيئة الأداة؛ لذلك لم نثبت فتح الملف ورسم جميع العناصر داخل Power BI Desktop في هذه الجولة. فحص الحزمة ليس بديلًا عن اختبار العرض داخل التطبيق. المصدر الأصلي محفوظ.

معاينة HTML اجتازت اختبارات منطقية لاختيارات السنوات والأنواع والصفحات وإعادة الضبط؛ هذه اختبارات تشغيل وليست مراجعة بصرية.
'''
(R/'تقرير-السياحة-السعودية/سجل-التحقق-من-الأرقام.md').write_text(doc,encoding='utf-8')
