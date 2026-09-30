from pathlib import Path
import openpyxl,json,math
ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'تقرير-السياحة-السعودية'
w=openpyxl.load_workbook(r'D:\السياحة\jmye_almnatq_2021_2025_3ea42fcd2d.xlsx',read_only=True,data_only=True)
data={'national':[],'regional':[],'monthly':{},'governorates':[]}
for key,table in [('national','Annual_National_Tourism'),('regional','Annual_Regional_Tourism')]:
    rows=json.loads((ROOT/'relationship-repair'/f'annual-{key}.json').read_text(encoding='utf-8-sig'))
    data[key]=[{k.split('[')[-1][:-1]:v for k,v in row.items()} for row in rows]
months=['January','February','March','April','May','June','July','August','September','October','November','December']
for sheet in ['D1','D2','D6','D7','D8']:
    year=None
    for row in w[sheet].iter_rows(values_only=True):
        if isinstance(row[1],(int,float)) and 2021<=row[1]<=2025:year=int(row[1])
        if year and row[2] in months:
            month=months.index(row[2])+1
            rec=data['monthly'].setdefault(f'{year}-{month:02}',{'year':year,'month':month})
            if sheet=='D1':rec.update(tourists=row[3]*1000,nights=row[4]*1000,spending=row[5]*1e6)
            elif sheet=='D2':rec['purpose']={k:row[i]*1000 for i,k in enumerate(['ديني','أعمال','ترفيه','أخرى','زيارة الأقارب والأصدقاء'],3)}
            elif sheet=='D6':rec['transport']={k:row[i]*1000 for i,k in enumerate(['جواً','براً','بحراً'],3)}
            elif sheet=='D7':rec['accommodation']={k:row[i]*1000 for i,k in enumerate(['الشقق','الفنادق','أخرى','السكن الخاص'],3)}
            else:rec['components']={k:row[i]*1e6 for i,k in enumerate(['الإيواء','الترفيه','الطعام والشراب','المواصلات','التسوق','أخرى'],3)}
s=w['D5'];headers=[c.value for c in s[10]];regions=[];region=None
for c in s[8]:
    if c.value is not None:region=c.value
    regions.append(region)
year=None;month=None;gov={}
for row in s.iter_rows(min_row=12,values_only=True):
    if row[1]=='Total':year=None;month=None;continue
    if isinstance(row[1],(int,float)) and 2021<=row[1]<=2025:year=int(row[1])
    if row[2] in months:month=months.index(row[2])+1
    if not year or not month or row[3] not in ['Tourists Number (Overnight Visitors)','Tourists Spending (SAR)']:continue
    field='tourists' if row[3].startswith('Tourists Number') else 'spending'
    for i in range(4,45):
        key=(year,month,headers[i]);rec=gov.setdefault(key,{'year':year,'month':month,'region':regions[i],'governorate':headers[i]})
        value=row[i];rec[field]=None if not isinstance(value,(int,float)) else value
data['governorates']=list(gov.values());data['monthly']=list(data['monthly'].values())
for y in range(2021,2026):
    a=sum(r['Tourists'] for r in data['national'] if r['Year']==y and r['TourismType']=='محلية')
    assert a==sum(r['tourists'] for r in data['monthly'] if r['year']==y)
    assert len([r for r in data['monthly'] if r['year']==y])==12
(ROOT/'report-build/data.json').write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
data['events']=json.loads((OUT/'سجل-الأحداث-والمصادر.json').read_text(encoding='utf-8'))
data['evidence']=json.loads((ROOT/'report-build/evidence-panels.json').read_text(encoding='utf-8'))
html=(ROOT/'report-build/report_template.html').read_text(encoding='utf-8').replace('__DATA__',json.dumps(data,ensure_ascii=False).replace('</','<\\/'))
(OUT/'معاينة-التقرير.html').write_text(html,encoding='utf-8')
print('Preview created with',len(data['monthly']),'months and',len(data['governorates']),'governorate-month records')
